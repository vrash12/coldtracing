from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\MAURICIO\Documents\coldtrace")
OUTPUT = ROOT / "output" / "docx" / "ColdTrace_Capstone_2_Progress_Report_1.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5F6B76"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9DEE5"
PALE_BLUE = "E8EEF5"
GREEN = "1B5E20"
PALE_GREEN = "E8F3EA"
WHITE = "FFFFFF"
BLACK = "000000"


MODULES = [
    (
        "Authentication and Role Access",
        "Active-account login, secure logout and session invalidation, role-based redirects, and server-side access checks for Administrator, Driver, and Receiver accounts.",
    ),
    (
        "Administrator Dashboard and Fleet Overview",
        "Operational counts for users, active devices, trips, and unresolved critical alerts, together with fleet and order monitoring based on the latest available telemetry.",
    ),
    (
        "User Account Management",
        "Administrator-only account create, view, edit, and delete functions; role and status assignment; search and filters; summary counts; and protection against deleting the signed-in account.",
    ),
    (
        "Administrator Order Management and Assignment",
        "Create, view, edit, cancel, and delete orders with multiple products, receiver and active-driver assignment, schedule conflict validation, trip synchronization, and driver assignment notification.",
    ),
    (
        "Receiver Order Portal",
        "Receiver dashboard, order search and status totals, multi-item order creation, ownership checks, pending-order editing or deletion, permitted cancellation, and saved delivery address and map location.",
    ),
    (
        "Driver Dashboard, Orders, and Notifications",
        "Driver-specific trip totals, assigned-order search and filters, current assignment details, unread assignment notifications, mark-one or mark-all read actions, and order-specific latest telemetry access.",
    ),
    (
        "Delivery Route Planning and Map Visualization",
        "Mapped delivery destinations, multi-stop route construction, route alternatives, ETA and distance display, optimized stop ordering, and browser location fallback when vehicle telemetry is unavailable.",
    ),
    (
        "Six-Device ESP32 Registry and MQTT Topics",
        "Server configuration and database seeding for ESP32-CT-1001 through ESP32-CT-1006, matching per-device telemetry topics, optional truck assignment, and a shared wildcard subscription topic.",
    ),
    (
        "Telemetry Intake and Device Heartbeat",
        "Validated POST /api/telemetry intake for configured devices; persistence of temperature, humidity, GPS, signal, and battery readings; device last-seen and active-status updates; and active-trip association.",
    ),
    (
        "Cold-Chain Analytics",
        "Mean Kinetic Temperature is calculated from recorded trip temperatures, Remaining Shelf Life is estimated for the assigned product, and both values are stored with telemetry for monitoring use.",
    ),
    (
        "Server-Authoritative Route Risk and Scoring",
        "PHP reloads the authenticated driver's order, trip, product, device, and latest database telemetry; verifies ownership and device-truck linkage; then computes temperature risk, RSL risk, and a normalized route score.",
    ),
    (
        "Validated AI Route Recommendation and Fallback",
        "The AI may select only a submitted route ID. Response structure, risk level, and confidence are validated, confidence is clamped to 0-1, and API, timeout, connection, or malformed-response failures return the deterministic lowest server score.",
    ),
    (
        "Administrative Reports and CSV Export",
        "Date-range statistics, order and trip status summaries, driver performance, daily temperature trends, temperature-breach counts, and CSV export for orders, trips, alerts, and telemetry.",
    ),
    (
        "Automated Verification",
        "The current suite passes 11 tests with 49 assertions, covering the six-device configuration, telemetry route and fleet rejection, server-owned route scoring, cross-truck telemetry rejection, AI response validation and fallback, and guest login redirection.",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_run(run, *, name="Calibri", size=11, color=BLACK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_table_paragraph(paragraph, *, size=9.2, color=BLACK, bold=False, align=None) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run(run, size=size, color=color, bold=bold)


def set_repeat_table_headers(table) -> None:
    repeat_header(table.rows[0])


def add_page_field(paragraph, field_code: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_node = OxmlElement("w:t")
    text_node.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    set_run(run, size=9, color=MUTED)
    for node in (begin, instr, separate, text_node, end):
        run._r.append(node)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("COLDTRACE | CAPSTONE 2 PROGRESS REPORT")
    set_run(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Page ")
    set_run(run, size=9, color=MUTED)
    add_page_field(p, "PAGE")
    run = p.add_run(" of ")
    set_run(run, size=9, color=MUTED)
    add_page_field(p, "NUMPAGES")


def add_metadata_table(doc: Document) -> None:
    rows = [
        ("Project Title", "ColdTrace: AI-Assisted Cold-Chain Monitoring and Delivery Management System"),
        ("Project Leader", "[Insert project leader name]"),
        ("Group Members", "[Insert group member names]"),
        ("Technical Adviser", "[Insert technical adviser name]"),
        ("Reporting Date", "August 15, 2026"),
    ]
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2160, 7200])
    for idx, text in enumerate(("Project Information", "Details")):
        cell = table.rows[0].cells[idx]
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        style_table_paragraph(
            cell.paragraphs[0],
            size=9.6,
            color=WHITE,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    repeat_header(table.rows[0])
    for idx, (label, value) in enumerate(rows, start=1):
        label_cell, value_cell = table.rows[idx].cells
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        value_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(label_cell, PALE_BLUE)
        p = label_cell.paragraphs[0]
        p.text = label
        style_table_paragraph(p, size=10, color=NAVY, bold=True)
        p = value_cell.paragraphs[0]
        p.text = value
        style_table_paragraph(p, size=10, color=BLACK)
        keep_row_together(table.rows[idx])


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(42)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("CAPSTONE 2")
    set_run(run, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("PROGRESS REPORT 1")
    set_run(run, size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run("ColdTrace: AI-Assisted Cold-Chain Monitoring\nand Delivery Management System")
    set_run(run, size=15, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run("Completed PHP/Laravel modules verified from the current project repository")
    set_run(run, size=10.5, color=MUTED, italic=True)

    add_metadata_table(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run("REPORT SCOPE")
    set_run(run, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(
        "This report lists only functions that are implemented in the current ColdTrace codebase. "
        "Incomplete or placeholder areas are intentionally omitted, and the listed scope is not an overall project-completion percentage."
    )
    set_run(run, size=10.5, color=BLACK)


def add_module_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["No.", "Completed Module", "Implemented Functionality", "Status"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        style_table_paragraph(
            cell.paragraphs[0],
            size=9.4,
            color=WHITE,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for number, (module, functionality) in enumerate(MODULES, start=1):
        cells = table.add_row().cells
        values = [str(number), module, functionality, "COMPLETED"]
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx == 3:
                set_cell_shading(cell, PALE_GREEN)
            elif number % 2 == 0:
                set_cell_shading(cell, "FAFBFC")

            align = WD_ALIGN_PARAGRAPH.LEFT
            bold = idx in (1, 3)
            color = GREEN if idx == 3 else BLACK
            size = 8.8 if idx == 2 else 9.0
            if idx in (0, 3):
                align = WD_ALIGN_PARAGRAPH.CENTER
            style_table_paragraph(
                cell.paragraphs[0],
                size=size,
                color=color,
                bold=bold,
                align=align,
            )
        keep_row_together(table.rows[-1])

    set_table_geometry(table, [648, 2520, 4968, 1224])
    set_repeat_table_headers(table)


def add_verification(doc: Document) -> None:
    heading = doc.add_paragraph("2. Verification Summary", style="Heading 1")
    heading.paragraph_format.page_break_before = False

    p = doc.add_paragraph()
    run = p.add_run(
        "Verification was performed using the project's existing automated test suite on August 15, 2026. "
        "The suite completed successfully with the following result:"
    )
    set_run(run, size=11)

    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Test result", "PASS"),
        ("Tests", "11 passed"),
        ("Assertions", "49 passed"),
        ("Device profiles verified", "ESP32-CT-1001 through ESP32-CT-1006"),
        ("Key safeguards verified", "Server-owned route scores, device-truck ownership, AI route validation, and deterministic fallback"),
    ]
    for idx, text in enumerate(("Verification Item", "Verified Result")):
        cell = table.rows[0].cells[idx]
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        style_table_paragraph(
            cell.paragraphs[0],
            size=9.4,
            color=WHITE,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    repeat_header(table.rows[0])
    for idx, (label, value) in enumerate(rows, start=1):
        left, right = table.rows[idx].cells
        left.text = label
        right.text = value
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(left, PALE_BLUE)
        if label == "Test result":
            set_cell_shading(right, PALE_GREEN)
        style_table_paragraph(left.paragraphs[0], size=9.6, color=NAVY, bold=True)
        style_table_paragraph(
            right.paragraphs[0],
            size=9.6,
            color=GREEN if label == "Test result" else BLACK,
            bold=label == "Test result",
        )
        keep_row_together(table.rows[idx])
    set_table_geometry(table, [2880, 6480])


def add_signoff(doc: Document) -> None:
    doc.add_paragraph("3. Review and Sign-off", style="Heading 1")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run(
        "The signatures below confirm review of the completed-module register for this reporting period."
    )
    set_run(run, size=10.5, color=MUTED)

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    labels = ["Prepared by", "Checked by"]
    details = [
        "[Project Leader / Group Representative]\n\nSignature: __________________________\nDate: ______________________________",
        "[Capstone 2 Subject Teacher]\n\nSignature: __________________________\nDate: ______________________________",
    ]
    for idx in range(2):
        header_cell = table.rows[0].cells[idx]
        header_cell.text = labels[idx]
        header_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(header_cell, NAVY)
        style_table_paragraph(
            header_cell.paragraphs[0],
            size=9.8,
            color=WHITE,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        cell = table.rows[1].cells[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_shading(cell, "FAFBFC")
        p = cell.paragraphs[0]
        p.text = details[idx]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.25
        for run in p.runs:
            set_run(run, size=9.6, color=BLACK)
    set_table_geometry(table, [4680, 4680])
    repeat_header(table.rows[0])
    keep_row_together(table.rows[0])
    keep_row_together(table.rows[1])


def build() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])

    doc.core_properties.title = "ColdTrace Capstone 2 Progress Report 1"
    doc.core_properties.subject = "Completed ColdTrace PHP/Laravel modules"
    doc.core_properties.author = "ColdTrace Project Team"
    doc.core_properties.keywords = "ColdTrace, capstone, progress report, cold chain, telemetry, route recommendation"

    add_cover(doc)
    doc.add_page_break()

    doc.add_paragraph("1. Completed System Functionality / Features / Modules", style="Heading 1")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(
        "The register below contains only modules and deliverables that are present in the current implementation. "
        "Each listed item is marked completed within this report's defined scope."
    )
    set_run(run, size=10.5)
    add_module_table(doc)

    add_verification(doc)
    add_signoff(doc)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
